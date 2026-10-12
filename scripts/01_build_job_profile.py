import json

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from project_config import (
    DOCS,
    DATA_PROCESSED,
    FIT_WEIGHTS,
    JOB_RIASEC_TARGET,
    PERSONALITY_DOMAINS,
    RIASEC,
)

JOB_PROFILE = {
    "role": "Project Manager - Digital Marketing Agency",
    "holland_code": "ESC",
    "targets": {
        "cognitive": {"min_percentile": 55},
        "personality": {
            "Neuroticism": [10, 40],
            "Extraversion": [40, 75],
            "Openness": [45, 80],
            "Agreeableness": [45, 85],
            "Conscientiousness": [65, 95],
        },
        "integrity": {"min_percentile": 65},
        "interests": JOB_RIASEC_TARGET,
    },
    "weights": FIT_WEIGHTS,
    "cv": {
        "education_weights": {
            "Marketing": 30,
            "Business Administration": 28,
            "Industrial Engineering": 26,
            "Advertising": 25,
            "Communications": 22,
            "Economics": 20,
            "Software Engineering": 22,
            "Psychology": 18,
            "Design": 15,
            "Other": 12,
        },
        "degree_bonus": {
            "Bachelor": 0,
            "Master": 5,
            "MBA": 6,
            "PhD": 7,
        },
        "certifications": [
            "PMP",
            "PRINCE2",
            "Certified ScrumMaster",
            "SAFe Agilist",
            "Google Analytics",
            "Meta Certified",
            "HubSpot Inbound",
            "Lean Six Sigma",
        ],
        "experience_ideal": [4, 12],
    },
}

EDUCATION_POOL = list(JOB_PROFILE["cv"]["education_weights"].keys())
DEGREES = list(JOB_PROFILE["cv"]["degree_bonus"].keys())
CERTIFICATIONS = JOB_PROFILE["cv"]["certifications"]


def _base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)


def _title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(17)
    return p


def _section(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    return p


def _bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p = doc.add_paragraph(style="List Bullet 2")
    p.add_run(text)
    return p


def build_job_profile():
    profile_path = DATA_PROCESSED / "job_profile.json"
    with open(profile_path, "w", encoding="utf-8") as fh:
        json.dump(JOB_PROFILE, fh, indent=2, ensure_ascii=False)
    return profile_path


def build_job_description_docx():
    doc = Document()
    _base_style(doc)

    _title(doc, "JOB DESCRIPTION")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Manager - Digital Marketing Agency")
    r.bold = True
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("Selection Process Reference Profile  |  Open-access psychometric route")

    _section(doc, "1. Job Summary")
    doc.add_paragraph(
        "The Project Manager owns the end-to-end delivery of digital marketing campaigns and client "
        "accounts inside a fast-growing digital agency. The role sits at the intersection of strategic "
        "planning, cross-functional coordination, and operational control: the right person turns "
        "ambiguous client goals into measurable, on-time, on-budget campaigns and keeps creative, "
        "media, content, developers and data teams aligned around a single plan."
    )

    _section(doc, "2. Key Responsibilities")
    for item in [
        "Own scope, timeline, budget and quality for a portfolio of 3-6 client projects simultaneously.",
        "Translate client business goals into campaign briefs, roadmaps and measurable KPIs.",
        "Coordinate creative, media, content, SEO, CRO and development squads on daily execution.",
        "Run agile ceremonies (sprint planning, stand-ups, reviews) and keep reporting cadences.",
        "Identify, assess and mitigate project risks; escalate exceptions with clear options and recommended decisions.",
        "Build trusted client relationships: written status updates, weekly meetings, change control and expectation management.",
        "Track budget burn, vendor invoices and profitability per account.",
        "Drive post-mortems and continuous improvement across the agency.",
    ]:
        _bullet(doc, item)

    _section(doc, "3. Technical Competencies")
    for item in [
        "Project planning and scheduling (WBS, Gantt, critical path).",
        "Agile / Scrum / Kanban delivery frameworks.",
        "Budgeting, forecasting and margin control.",
        "Digital marketing fundamentals: performance media, content, SEO, social, email, analytics.",
        "Reporting and data literacy: campaign KPIs, dashboards and executive summaries.",
        "Tooling: project platforms (Asana, ClickUp, Jira), G-Suite, Google Analytics, media buying interfaces.",
    ]:
        _bullet(doc, item)

    _section(doc, "4. Behavioral Competencies")
    for item in [
        "Structured, organized and self-disciplined execution under multiple deadlines.",
        "Assertive communication and stakeholder management, including difficult conversations.",
        "Emotional control and composure under pressure and ambiguity.",
        "Collaboration, empathy and team dynamics awareness.",
        "Commercial mindset: curiosity, integrity and risk-awareness in decisions.",
    ]:
        _bullet(doc, item)

    _section(doc, "5. Requirements")
    for item in [
        "Bachelor's degree in Marketing, Business Administration, Industrial Engineering, Advertising, Communications or related field (Master's / MBA valued).",
        "4-12 years of professional experience, with at least 2 years managing digital or tech projects.",
        "Prior experience in a digital marketing agency or in-house marketing operations preferred.",
        "Certifications valued: PMP, PRINCE2, Certified ScrumMaster, SAFe Agilist, Google Analytics, Meta Certified, HubSpot Inbound, Lean Six Sigma.",
    ]:
        _bullet(doc, item)

    _section(doc, "6. Psychometric Target Profile (anchor for selection)")
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Instrument"
    hdr[1].text = "Construct measured"
    hdr[2].text = "Target zone (percentile)"
    rows = [
        ("ICAR (16-item)", "Cognitive ability", ">= 55 (medium-high)"),
        ("IPIP-NEO-120", "Personality - Conscientiousness", "65 - 95"),
        ("IPIP-NEO-120", "Personality - Extraversion", "40 - 75"),
        ("IPIP-NEO-120", "Personality - Openness", "45 - 80"),
        ("IPIP-NEO-120", "Personality - Agreeableness", "45 - 85"),
        ("IPIP-NEO-120", "Personality - Neuroticism", "10 - 40"),
        ("IPIP-HEXACO H-H", "Integrity (Honesty-Humility)", ">= 65"),
        ("O*NET Interest Profiler", "Vocational interests (RIASEC)", "ESC-dominant profile"),
    ]
    for k, v, z in rows:
        row = tbl.add_row().cells
        row[0].text = k
        row[1].text = v
        row[2].text = z

    _section(doc, "7. What Success Looks Like in the First 90 Days")
    for item in [
        "Take ownership of an assigned portfolio and deliver a first client roadmap within the first two weeks.",
        "Stabilize reporting: every active account has a live KPI dashboard and a weekly client rhythm.",
        "Improve on-time delivery and margin vs. the agency baseline.",
        "Be identified by clients and squads as the reference point for project status and decisions.",
    ]:
        _bullet(doc, item)

    _section(doc, "8. Selection Process")
    doc.add_paragraph(
        "Candidates passed through a standardized multi-phase process: document and CV review, "
        "psychometric assessment (cognitive, personality, integrity and vocational interests), "
        "psychological interview, technical/labor interview and an assessment center with behavioral "
        "and case exercises. Results are consolidated in this dashboard."
    )

    out = DOCS / "job_description_PM.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    j = build_job_profile()
    d = build_job_description_docx()
    print("job profile:", j)
    print("docx:", d)